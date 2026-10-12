"""
Export Word (python-docx) à partir des tokens du Document.

Deux modes
----------
classique : mapping direct vers les styles Word standard (titres, listes,
            tableaux "Table Grid", images à leur taille d'origine plafonnée
            à la largeur utile). Rapide, sobre, fidèle au contenu.
avancé    : en plus — styles redéfinis (police, titres), table des matières
            et numéros de page (champs Word), callouts Obsidian en encadrés
            colorés, code avec coloration syntaxique et fond, tableaux avec
            en-tête répété, légendes d'images, surlignage, citations avec
            filet, métadonnées, page de titre, et optimisation des images
            (rééchantillonnage au DPI cible + compression JPEG).

Qualité des images : `quality` = brouillon | standard | haute | originale.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageOps

from .core import Document, callout_style

try:
    from pygments import lex
    from pygments.lexers import TextLexer, get_lexer_by_name
    from pygments.styles import get_style_by_name
    from pygments.util import ClassNotFound
except ImportError:
    lex = None

QUALITY_PRESETS = {
    "brouillon": dict(dpi=96, jpeg=70),
    "standard": dict(dpi=150, jpeg=85),
    "haute": dict(dpi=220, jpeg=92),
    "originale": dict(dpi=0, jpeg=100),   # dpi=0 : aucune modification de l'image
}


@dataclass
class ExportOptions:
    mode: str = "advanced"            # "classic" | "advanced"
    quality: str = "standard"         # clé de QUALITY_PRESETS
    image_dpi: int | None = None      # surcharge du preset
    jpeg_quality: int | None = None   # surcharge du preset
    image_format: str = "auto"        # "auto" | "png" | "jpeg"
    image_max_width_cm: float | None = None   # None = largeur utile de la page
    page_size: str = "A4"             # "A4" | "Letter"
    margin_cm: float = 2.2
    template: Path | None = None      # .docx/.dotx fournissant les styles
    font: str = "Calibri"
    font_size: float = 11
    code_font: str = "Consolas"
    code_font_size: float = 9.5
    # --- avancé uniquement ---
    toc: bool = True
    toc_depth: int = 3
    title_page: bool = False
    page_numbers: bool = True
    header_text: str = ""
    callouts: bool = True
    captions: bool = True
    syntax_highlight: bool = True

    @property
    def dpi(self) -> int:
        return QUALITY_PRESETS[self.quality]["dpi"] if self.image_dpi is None else self.image_dpi

    @property
    def jpeg(self) -> int:
        return QUALITY_PRESETS[self.quality]["jpeg"] if self.jpeg_quality is None else self.jpeg_quality


# ---------------------------------------------------------------------------
# Helpers XML (l'ordre des éléments est imposé par le schéma OOXML)
# ---------------------------------------------------------------------------

_P_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
                "w:sectPr", "w:pPrChange")
_P_AFTER_BDR = ("w:shd",) + _P_AFTER_SHD
_R_AFTER_SHD = ("w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
                "w:eastAsianLayout", "w:specVanish", "w:oMath")
_TC_AFTER_BDR = ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                 "w:vAlign", "w:hideMark")
_TC_AFTER_SHD = _TC_AFTER_BDR[1:]


def _shd(fill: str) -> OxmlElement:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), fill)
    return el


def _shade_ppr(ppr, fill):
    ppr.insert_element_before(_shd(fill), *_P_AFTER_SHD)


def _shade_run(run, fill):
    run._r.get_or_add_rPr().insert_element_before(_shd(fill), *_R_AFTER_SHD)


def _shade_cell(cell, fill):
    cell._tc.get_or_add_tcPr().insert_element_before(_shd(fill), *_TC_AFTER_SHD)


def _cell_left_border(cell, color, sz=24):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        if side == "left":
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcpr.insert_element_before(borders, *_TC_AFTER_BDR)


def _p_border(p, side, color="999999", sz=12, space=6):
    ppr = p._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(bdr, *_P_AFTER_BDR)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color)
    bdr.append(el)


def _field(p, instr, placeholder=""):
    """Champ Word complexe (PAGE, NUMPAGES, TOC...)."""
    def fld(kind):
        r = p.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
    fld("begin")
    r = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fld("separate")
    p.add_run(placeholder)
    fld("end")


def _hyperlink(p, run, url):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    h.set(qn("w:history"), "1")
    run._r.addprevious(h)
    h.append(run._r)


def _set_style_font(style, name):
    style.font.name = name
    rfonts = style.element.get_or_add_rPr().rFonts
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{a}"), None)   # sinon le thème écrase la police
    rfonts.set(qn("w:eastAsia"), name)


def _close(toks, i):
    level = 0
    for j in range(i, len(toks)):
        level += toks[j].nesting
        if level == 0:
            return j
    return len(toks) - 1


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------

def prepare_image(path: Path, width_cm: float, o: ExportOptions, advanced: bool) -> io.BytesIO:
    """Retourne les octets à embarquer (PNG/JPEG) pour une image affichée à width_cm."""
    suffix = path.suffix.lower()
    dpi = o.dpi

    if suffix == ".svg":                       # Word ne gère pas le SVG de façon fiable
        import cairosvg
        px = max(1, int(width_cm / 2.54 * (dpi or 220 if advanced else 150)))
        return io.BytesIO(cairosvg.svg2png(url=str(path), output_width=px))

    if (not advanced or dpi == 0) and suffix in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"):
        return io.BytesIO(path.read_bytes())   # classique / originale : aucun traitement

    img = ImageOps.exif_transpose(Image.open(path))
    img.load()
    has_alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)

    resized = False
    if dpi:
        target = max(1, int(width_cm / 2.54 * dpi))
        if img.width > target:
            img = img.resize((target, round(img.height * target / img.width)), Image.LANCZOS)
            resized = True

    fmt = o.image_format
    if fmt == "auto":
        fmt = "jpeg" if suffix in (".jpg", ".jpeg") else "png"
    if has_alpha:
        fmt = "png"

    # Rien à gagner : on évite une recompression JPEG inutile
    if not resized and ((fmt == "png" and suffix == ".png") or (fmt == "jpeg" and suffix in (".jpg", ".jpeg"))):
        return io.BytesIO(path.read_bytes())

    buf = io.BytesIO()
    if fmt == "png":
        img.convert("RGBA" if has_alpha else "RGB").save(buf, "PNG", optimize=True)
    else:
        img.convert("RGB").save(buf, "JPEG", quality=o.jpeg, optimize=True)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Exporteur
# ---------------------------------------------------------------------------

class DocxExporter:
    def __init__(self, options: ExportOptions | None = None):
        self.o = options or ExportOptions()
        self.adv = self.o.mode == "advanced"

    # -- point d'entrée ----------------------------------------------------

    def export(self, doc: Document, out_path) -> Path:
        self.src = doc
        self.d = DocxDocument(str(self.o.template)) if self.o.template else DocxDocument()
        self.container = self.d
        self.list_level = 0
        self.quote_level = 0
        self._marker = None

        if not self.o.template:
            self._setup_page()
            self._setup_styles()
        sec = self.d.sections[0]
        self._usable_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
        self._avail_cm = self._usable_cm

        cp = self.d.core_properties
        cp.title = doc.title
        cp.author = doc.meta.get("author", "")

        if self.adv:
            self._header_footer()
            if self.o.title_page:
                self._title_page()
            if self.o.toc:
                self._toc()

        self._blocks(doc.tokens)

        if self.adv and self.o.toc:    # Word propose de mettre à jour la TOC à l'ouverture
            uf = OxmlElement("w:updateFields")
            uf.set(qn("w:val"), "true")
            self.d.settings.element.append(uf)

        out_path = Path(out_path)
        self.d.save(str(out_path))
        return out_path

    # -- mise en page / styles --------------------------------------------

    def _setup_page(self):
        sec = self.d.sections[0]
        if self.o.page_size.lower() == "letter":
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        else:
            sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(self.o.margin_cm))

    def _setup_styles(self):
        o, st = self.o, self.d.styles
        normal = st["Normal"]
        _set_style_font(normal, o.font)
        normal.font.size = Pt(o.font_size)
        normal.paragraph_format.space_after = Pt(6)

        if self.adv:
            for n, size in zip(range(1, 7), (20, 16, 14, 12, 11, 11)):
                s = st[f"Heading {n}"]
                _set_style_font(s, o.font)
                s.font.size = Pt(size)
                s.font.bold = True
                s.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
                s.paragraph_format.space_before = Pt(14 if n < 3 else 10)
                s.paragraph_format.space_after = Pt(6)
                s.paragraph_format.keep_with_next = True

        try:
            code = st["Code"]
        except KeyError:
            code = st.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        _set_style_font(code, o.code_font)
        code.font.size = Pt(o.code_font_size)
        pf = code.paragraph_format
        pf.space_before = pf.space_after = Pt(4)
        pf.line_spacing = 1.0
        if self.adv:
            _shade_ppr(code.element.get_or_add_pPr(), "F6F8FA")

    def _header_footer(self):
        sec = self.d.sections[0]
        if self.o.header_text:
            hp = sec.header.paragraphs[0]
            hp.text = self.o.header_text
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if self.o.page_numbers:
            fp = sec.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fp.add_run("Page ")
            _field(fp, "PAGE", "1")
            fp.add_run(" / ")
            _field(fp, "NUMPAGES", "1")

    def _title_page(self):
        for _ in range(6):
            self.d.add_paragraph()
        self._par("Title").add_run(self.src.title)
        sub = " — ".join(x for x in (self.src.meta.get("author"), self.src.meta.get("date")) if x)
        if sub:
            self._par("Subtitle").add_run(sub)
        self.d.add_page_break()

    def _toc(self):
        p = self.d.add_paragraph()
        r = p.add_run("Table des matières")
        r.bold = True
        r.font.size = Pt(16)
        p = self.d.add_paragraph()
        _field(p, f'TOC \\o "1-{self.o.toc_depth}" \\h \\z \\u',
               "Clic droit › Mettre à jour les champs pour générer la table des matières.")
        self.d.add_page_break()

    # -- paragraphes -------------------------------------------------------

    def _par(self, style=None):
        try:
            return self.container.add_paragraph(style=style)
        except KeyError:                       # style absent d'un modèle personnalisé
            return self.container.add_paragraph()

    def _spacer(self):
        p = self.container.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(6)

    def _start_par(self):
        """Crée un paragraphe en tenant compte du contexte (liste, citation)."""
        marker, self._marker = self._marker, None
        if marker and marker[0] == "bullet":
            level = min(marker[1], 3)
            return self._par("List Bullet" + ("" if level == 1 else f" {level}"))

        if marker:                              # liste numérotée : numéro écrit en dur
            _, level, num = marker
            p = self._par("List Paragraph")
            indent = Cm(0.9 * level)
            p.paragraph_format.left_indent = indent
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.tab_stops.add_tab_stop(indent)
            p.add_run(f"{num}.\t")
            return p

        if self.list_level:                     # paragraphe de suite d'un item
            p = self._par()
            p.paragraph_format.left_indent = Cm(0.9 * self.list_level)
            return p

        if self.quote_level:
            p = self._par("Quote")
            p.paragraph_format.left_indent = Cm(0.8 * self.quote_level)
            if self.adv:
                _p_border(p, "left", "9AA5B1", sz=18, space=8)
            return p
        return self._par()

    # -- blocs -------------------------------------------------------------

    def _blocks(self, toks):
        i = 0
        while i < len(toks):
            handler = getattr(self, f"_b_{toks[i].type}", None)
            i = handler(toks, i) if handler else i + 1

    def _b_heading_open(self, toks, i):
        level = int(toks[i].tag[1])
        p = self._par(f"Heading {min(level, 9)}")
        self._inline(p, toks[i + 1], block_ok=False)
        return i + 3

    def _b_paragraph_open(self, toks, i):
        inline = toks[i + 1]
        if not inline.children:
            return i + 3
        p = self._start_par()
        self._inline(p, inline)
        return i + 3

    def _b_hr(self, toks, i):
        p = self._par()
        if self.adv:
            _p_border(p, "bottom")
        else:
            p.add_run("―" * 30)
        return i + 1

    def _b_html_block(self, toks, i):
        text = re.sub(r"<[^>]+>", "", toks[i].content).strip()
        if text:
            self._par().add_run(text)
        return i + 1

    def _b_bullet_list_open(self, toks, i):
        return self._list(toks, i, ordered=False)

    def _b_ordered_list_open(self, toks, i):
        return self._list(toks, i, ordered=True)

    def _list(self, toks, i, ordered):
        end = _close(toks, i)
        n = int(toks[i].attrGet("start") or 1)
        self.list_level += 1
        k = i + 1
        while k < end:
            if toks[k].type == "list_item_open":
                ke = _close(toks, k)
                self._marker = ("number" if ordered else "bullet", self.list_level, n)
                self._blocks(toks[k + 1:ke])
                self._marker = None
                n += 1
                k = ke + 1
            else:
                k += 1
        self.list_level -= 1
        return end + 1

    def _b_blockquote_open(self, toks, i):
        end = _close(toks, i)
        callout = toks[i].meta.get("callout")
        if callout and self.adv and self.o.callouts:
            self._callout(callout, toks[i + 1:end])
        else:
            self.quote_level += 1
            if callout:
                self._par("Quote").add_run(callout["title"]).bold = True
            self._blocks(toks[i + 1:end])
            self.quote_level -= 1
        return end + 1

    def _callout(self, c, inner):
        icon, color, bg = callout_style(c["type"])
        table = self.container.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        _shade_cell(cell, bg)
        _cell_left_border(cell, color)

        saved = (self.container, self._avail_cm, self.list_level, self.quote_level)
        self.container, self._avail_cm = cell, self._avail_cm - 1.0
        self.list_level = self.quote_level = 0

        head = cell.paragraphs[0]
        head.paragraph_format.space_after = Pt(3)
        r = head.add_run(f"{icon} {c['title']}")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(color)
        self._blocks(inner)

        self.container, self._avail_cm, self.list_level, self.quote_level = saved
        self._spacer()

    def _b_fence(self, toks, i):
        tok = toks[i]
        code = tok.content.rstrip("\n")
        lang = (tok.info or "").strip().split(" ")[0]
        p = self._par("Code")
        if self.adv and self.o.syntax_highlight and lex is not None:
            try:
                lexer = get_lexer_by_name(lang) if lang else TextLexer()
            except ClassNotFound:
                lexer = TextLexer()
            style = get_style_by_name("default")
            parts = [(v, style.style_for_token(t)) for t, v in lex(code, lexer)]
            if parts:                                  # pygments ajoute un \n final
                parts[-1] = (parts[-1][0].rstrip("\n"), parts[-1][1])
            for value, st in parts:
                self._code_text(p, value, st["color"], st["bold"], st["italic"])
        else:
            self._code_text(p, code)
        return i + 1

    _b_code_block = _b_fence

    @staticmethod
    def _code_text(p, text, color=None, bold=False, italic=False):
        for n, part in enumerate(text.split("\n")):
            if n:
                p.add_run().add_break()
            if part:
                r = p.add_run(part)
                if color:
                    r.font.color.rgb = RGBColor.from_string(color.upper())
                r.bold = True if bold else None
                r.italic = True if italic else None

    def _b_table_open(self, toks, i):
        end = _close(toks, i)
        rows, cur, pending = [], None, (None, False)
        for t in toks[i:end]:
            if t.type == "tr_open":
                cur = []
            elif t.type == "tr_close":
                rows.append(cur)
            elif t.type in ("th_open", "td_open"):
                style = str(t.attrGet("style") or "")
                align = re.search(r"text-align:\s*(\w+)", style)
                pending = (align.group(1) if align else None, t.type == "th_open")
            elif t.type == "inline":
                cur.append((t, *pending))
        if not rows:
            return end + 1

        ncols = max(len(r) for r in rows)
        table = self.container.add_table(rows=len(rows), cols=ncols)
        try:
            table.style = "Table Grid"
        except Exception:
            pass

        saved = self._avail_cm
        self._avail_cm = saved / ncols - 0.4
        aligns = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                  "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for r, row in enumerate(rows):
            for c, (inline, align, header) in enumerate(row):
                cell = table.cell(r, c)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                if align in aligns:
                    p.alignment = aligns[align]
                self._inline(p, inline, block_ok=False)
                if header:
                    for run in p.runs:
                        run.bold = True
                    if self.adv:
                        _shade_cell(cell, "DCE6F2")
            if self.adv and r == 0 and any(h for *_, h in row):
                trpr = table.rows[0]._tr.get_or_add_trPr()
                el = OxmlElement("w:tblHeader")
                el.set(qn("w:val"), "true")
                trpr.append(el)
        self._avail_cm = saved
        self._spacer()
        return end + 1

    # -- contenu en ligne ---------------------------------------------------

    def _inline(self, p, inline, block_ok=True):
        kids = inline.children or []
        if len(kids) == 1 and kids[0].type == "image" and "path" in kids[0].meta:
            self._image(p, kids[0], block=block_ok)
            return

        state = {"strong": 0, "em": 0, "s": 0, "mark": 0}
        link = None
        for k in kids:
            t = k.type
            if t == "text":
                self._run(p, k.content, state, link)
            elif t == "code_inline":
                self._run(p, k.content, state, link, code=True)
            elif t.endswith("_open") and t[:-5] in state:
                state[t[:-5]] += 1
            elif t.endswith("_close") and t[:-6] in state:
                state[t[:-6]] -= 1
            elif t == "link_open":
                link = str(k.attrGet("href") or "")
            elif t == "link_close":
                link = None
            elif t == "softbreak":
                if self.src.breaks:
                    p.add_run().add_break()
                else:
                    p.add_run(" ")
            elif t == "hardbreak":
                p.add_run().add_break()
            elif t == "image":
                self._image(p, k, block=False)
            elif t == "html_inline" and re.match(r"<br\s*/?>", k.content, re.I):
                p.add_run().add_break()

    def _run(self, p, text, state, link, code=False):
        r = p.add_run(text)
        r.bold = True if state["strong"] else None
        r.italic = True if state["em"] else None
        r.font.strike = True if state["s"] else None
        if state["mark"]:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if code:
            r.font.name = self.o.code_font
            r.font.size = Pt(self.o.code_font_size)
            if self.adv:
                _shade_run(r, "EEF0F3")
        if link:
            r.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            if not link.startswith("wikilink:"):
                r.font.underline = True
                _hyperlink(p, r, link)

    def _image(self, p, tok, block):
        meta = tok.meta
        max_cm = min(self.o.image_max_width_cm or self._usable_cm, self._avail_cm)
        width_cm = min(meta["width_px"] / 96 * 2.54, max_cm)       # 96 dpi = convention Obsidian/CSS
        data = prepare_image(Path(meta["path"]), width_cm, self.o, self.adv)
        shape = p.add_run().add_picture(data, width=Cm(width_cm))

        alt = tok.content or ""
        shape._inline.docPr.set("descr", alt)                      # texte alternatif (accessibilité)
        if block:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if self.adv and self.o.captions and alt:
                p.paragraph_format.keep_with_next = True
                cap = self._par("Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.add_run(alt)


def export_docx(doc: Document, out_path, options: ExportOptions | None = None) -> Path:
    return DocxExporter(options).export(doc, out_path)
